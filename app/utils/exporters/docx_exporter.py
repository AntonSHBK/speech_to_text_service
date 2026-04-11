from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.utils.exporters.common import build_paragraph_blocks, format_timestamp


def _configure_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph_format = paragraph.paragraph_format
    paragraph_format.line_spacing = 1.25
    paragraph_format.space_after = Pt(6)
    paragraph_format.space_before = Pt(6)


def _part_to_chunk(part: dict, export_timestamps: bool) -> str:
    text = str(part.get("text", "")).strip()
    if not text:
        return ""
    if not export_timestamps:
        return text
    start = format_timestamp(part.get("start"))
    end = format_timestamp(part.get("end"))
    return f"[{start} - {end}] {text}"


def _append_block_paragraph(
    document: Document,
    block: dict,
    export_timestamps: bool,
    show_speaker: bool = True,
) -> None:
    parts = block.get("parts") or []
    speaker = block.get("speaker")

    paragraph = document.add_paragraph()
    _configure_paragraph(paragraph)

    if speaker and show_speaker:
        speaker_run = paragraph.add_run(f"{speaker}: ")
        speaker_run.bold = True

    first_chunk = True
    for part in parts:
        chunk = _part_to_chunk(part, export_timestamps=export_timestamps)
        if not chunk:
            continue
        if not first_chunk:
            paragraph.add_run(" ")
        paragraph.add_run(chunk)
        first_chunk = False


def export_docx(
    result: dict,
    path: Path,
    export_timestamps: bool = False,
    paragraph_pause_sec: float = 2.0,
    paragraph_max_chars: int = 350,
) -> Path:
    document = Document()

    heading = document.add_heading("Результат транскрибации", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    blocks = build_paragraph_blocks(
        result,
        pause_sec=paragraph_pause_sec,
        max_chars=paragraph_max_chars,
    )
    if blocks:
        prev_speaker: str | None = None
        for block in blocks:
            speaker = block.get("speaker")
            _append_block_paragraph(
                document,
                block,
                export_timestamps=export_timestamps,
                show_speaker=(speaker != prev_speaker),
            )
            prev_speaker = speaker if isinstance(speaker, str) else None
    else:
        paragraph = document.add_paragraph(result.get("text", "") or "")
        _configure_paragraph(paragraph)

    document.save(path)
    return path
